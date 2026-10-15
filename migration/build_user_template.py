"""Build the Excel templates admins use to bulk-import users.

Usage:  python build_user_template.py
Writes  import/user_import_template.xlsx   blank template (what the backoffice offers for download)
        import/province_managers_76.xlsx   same template with one manager row per province (only if missing)
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
provinces = sorted(json.load(open(os.path.join(HERE, 'import', 'provinces.json'), encoding='utf-8')), key=lambda p: p['name_th'])
ROLES = ['ผู้ดูแลจังหวัด', 'เจ้าหน้าที่จังหวัด', 'ผู้ดูแลระบบส่วนกลาง']
COLUMNS = [  # header, width, required, help  (rules match web/src/lib/fieldRules.ts)
    ('ชื่อผู้ใช้ (username)', 24, True, 'a-z ตัวพิมพ์เล็ก ตัวเลข จุด (.) ขีด (- _) ขึ้นต้นด้วยตัวอักษรหรือตัวเลข ยาว 4-32 ตัว ห้ามซ้ำ เช่น somchai.s'),
    ('ชื่อ นามสกุล', 30, True, 'ตัวอักษรไทยหรืออังกฤษ ไม่มีตัวเลข เว้นวรรคระหว่างชื่อและนามสกุล ใส่คำนำหน้าหรือไม่ก็ได้'),
    ('ตำแหน่ง', 28, False, 'ไม่บังคับ ตัวอักษรไทย อังกฤษ ตัวเลข และ . , ( ) / - & (ไม่ใช่ตัวเลขล้วน)'),
    ('อีเมล', 32, False, 'ไม่บังคับ รูปแบบอีเมล เช่น somchai@nso.go.th ใช้ร่างอีเมลส่งลิงก์เปิดใช้บัญชี'),
    ('โทรศัพท์', 18, False, 'ไม่บังคับ 9-10 หลัก ขึ้นต้นด้วย 0 ใส่ขีดหรือเว้นวรรคได้ เช่น 081-234-5678'),
    ('บทบาท', 24, True, 'เลือกจากรายการ'),
    ('จังหวัด', 22, True, 'เลือกจากรายการ; เว้นว่างได้เฉพาะบทบาทผู้ดูแลระบบส่วนกลาง'),
    ('ชื่อ-นามสกุล (ภาษาอังกฤษ)', 30, False, 'ไม่บังคับ ตัวอักษรอังกฤษ ชื่อและนามสกุลเว้นวรรคกัน เช่น Somchai Jaidee'),
]
MAX_ROWS = 500
HEAD_FILL = PatternFill('solid', fgColor='083F58')
REQ_FILL = PatternFill('solid', fgColor='0B5A7D')
THIN = Side(style='thin', color='C5D0D9')
FONT = 'TH Sarabun New'


def build(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'ผู้ใช้'
    for c, (title, width, required, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=title + (' *' if required else ''))
        cell.font = Font(name=FONT, size=16, bold=True, color='FFFFFF')
        cell.fill = REQ_FILL if required else HEAD_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = Border(bottom=THIN, right=THIN)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            ws.cell(row=r, column=c, value=value)
    for r in range(2, MAX_ROWS + 2):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=16)
            cell.number_format = '@'          # keep phone numbers and usernames as text
    role_dv = DataValidation(type='list', formula1='"' + ','.join(ROLES) + '"', allow_blank=False,
                             showErrorMessage=True, errorTitle='บทบาทไม่ถูกต้อง', error='เลือกบทบาทจากรายการ')
    prov_dv = DataValidation(type='list', formula1=f"='รายชื่อจังหวัด'!$A$2:$A${len(provinces) + 1}", allow_blank=True,
                             showErrorMessage=True, errorTitle='จังหวัดไม่ถูกต้อง', error='เลือกจังหวัดจากรายการ')
    ws.add_data_validation(role_dv)
    ws.add_data_validation(prov_dv)
    role_dv.add(f'F2:F{MAX_ROWS + 1}')
    prov_dv.add(f'G2:G{MAX_ROWS + 1}')

    info = wb.create_sheet('คำอธิบาย')
    lines = [('วิธีใช้', None),
             ('1. กรอกข้อมูลในแผ่น "ผู้ใช้" แถวละ 1 คน เริ่มที่แถว 2 ห้ามแก้หรือสลับหัวคอลัมน์', None),
             ('2. คอลัมน์ที่มี * ต้องกรอก', None),
             ('3. บันทึกเป็นไฟล์ .xlsx แล้วอัปโหลดที่ หลังบ้าน > ผู้ใช้ > นำเข้าจากไฟล์', None),
             ('4. ระบบจะตรวจทุกแถวและแสดงผลให้ดูก่อน ยังไม่สร้างบัญชีจนกว่าจะกดยืนยัน', None),
             ('5. เมื่อยืนยัน ระบบสร้างลิงก์เปิดใช้บัญชีแบบใช้ครั้งเดียวให้แต่ละคน ดาวน์โหลดเป็นไฟล์แล้วส่งให้เจ้าตัวทีละคน', None),
             ('   เจ้าตัวเปิดลิงก์เพื่อตั้งรหัสผ่านและตั้งค่า 2FA เอง ไฟล์นี้จึงไม่มีช่องรหัสผ่าน', None),
             (f'6. นำเข้าได้ครั้งละไม่เกิน {MAX_ROWS} คน', None),
             ('7. ระบบตรวจรูปแบบข้อมูลทุกช่องตามคำอธิบายด้านล่าง แถวที่ผิดจะแสดงเหตุผลรายช่อง และไม่ถูกสร้าง', None),
             ('', None), ('คอลัมน์', 'คำอธิบาย')]
    lines += [(t + (' *' if req else ''), h) for t, _, req, h in COLUMNS]
    for r, (a, b) in enumerate(lines, 1):
        info.cell(row=r, column=1, value=a).font = Font(name=FONT, size=16, bold=(b is None and r == 1) or a in ('คอลัมน์',))
        if b:
            info.cell(row=r, column=2, value=b).font = Font(name=FONT, size=16, bold=a == 'คอลัมน์')
    info.column_dimensions['A'].width = 34
    info.column_dimensions['B'].width = 90

    ref = wb.create_sheet('รายชื่อจังหวัด')
    for c, t in enumerate(['จังหวัด', 'รหัส'], 1):
        ref.cell(row=1, column=c, value=t).font = Font(name=FONT, size=16, bold=True)
    for r, p in enumerate(provinces, 2):
        ref.cell(row=r, column=1, value=p['name_th']).font = Font(name=FONT, size=16)
        ref.cell(row=r, column=2, value=p['code']).font = Font(name=FONT, size=16)
    ref.column_dimensions['A'].width = 24
    ref.protection.sheet = True
    wb.save(path)


build(os.path.join(HERE, 'import', 'user_import_template.xlsx'), [])
# The managers file is filled in (admin_{slug}, office name) and checked: only create it when it is missing,
# never write over it (rebuilding the blank template once wiped the filled list).
managers = os.path.join(HERE, 'import', 'province_managers_76.xlsx')
if not os.path.exists(managers):
    build(managers, [(None, None, None, None, None, ROLES[0], p['name_th']) for p in provinces])
print('ok', len(provinces))
